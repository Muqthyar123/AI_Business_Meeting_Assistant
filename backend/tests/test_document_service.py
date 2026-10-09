"""Tests for DocumentExtractionService handling PDF, DOCX, TXT, and normalization."""

import io
from unittest.mock import MagicMock, patch
import docx
import pytest
from app.schemas.document import DocumentSourceType
from app.services.document_service import (
    CorruptDocumentError,
    DocumentExtractionService,
    EmptyDocumentError,
    NoExtractableTextError,
    UnsupportedDocumentFormatError,
    get_document_extraction_service,
)


@pytest.fixture
def service():
    """Document extraction service instance."""
    return DocumentExtractionService()


def test_normalize_text():
    """Verify whitespace normalization preserves paragraphs and strips excess spaces."""
    raw = "  Line   one   with   spaces.\r\n\r\n\r\n\r\nLine two with\t\ttabs.  \n\nLine three.  "
    normalized = DocumentExtractionService.normalize_text(raw)
    assert "Line one with spaces." in normalized
    assert "Line two with tabs." in normalized
    assert "Line three." in normalized
    # Max 2 newlines (paragraph separation)
    assert "\n\n\n" not in normalized
    assert not normalized.startswith(" ")
    assert not normalized.endswith(" ")


# --- TXT Extraction Tests ---


def test_extract_from_txt_utf8(service):
    """Verify TXT extraction with UTF-8 encoding."""
    content = "Meeting Title: Strategy Sync\n\nKey discussion items:\n- Budget approval\n- Hiring goals"
    text = service.extract_from_txt(content.encode("utf-8"))
    assert "Meeting Title: Strategy Sync" in text
    assert "- Budget approval" in text


def test_extract_from_txt_latin1(service):
    """Verify TXT extraction with Latin-1 fallback encoding."""
    content = "Meeting in Café Zürich: Résumé of actions"
    text = service.extract_from_txt(content.encode("latin-1"))
    assert "Café Zürich" in text


def test_extract_from_txt_empty(service):
    """Verify empty TXT raises EmptyDocumentError."""
    with pytest.raises(EmptyDocumentError, match="empty"):
        service.extract_from_txt(b"")

    with pytest.raises(EmptyDocumentError, match="no extractable text"):
        service.extract_from_txt(b"   \n\n\t  ")


# --- DOCX Extraction Tests ---


def test_extract_from_docx_valid(service):
    """Verify extraction from a valid DOCX file containing paragraphs and tables."""
    doc = docx.Document()
    doc.add_heading("Sprint Retro", level=1)
    doc.add_paragraph("What went well: Deployment was smooth.")
    doc.add_paragraph("Action item: Improve test coverage.")

    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Task"
    table.rows[0].cells[1].text = "Owner"
    table.rows[1].cells[0].text = "Write docs"
    table.rows[1].cells[1].text = "Alex"

    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    extracted = service.extract_from_docx(docx_bytes)
    assert "Sprint Retro" in extracted
    assert "What went well: Deployment was smooth." in extracted
    assert "Write docs | Alex" in extracted


def test_extract_from_docx_empty(service):
    """Verify empty or whitespace-only DOCX raises EmptyDocumentError."""
    doc = docx.Document()
    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    with pytest.raises(EmptyDocumentError, match="no extractable text"):
        service.extract_from_docx(docx_bytes)


def test_extract_from_docx_corrupt(service):
    """Verify corrupted DOCX raises CorruptDocumentError."""
    corrupt_bytes = b"PK\x03\x04corrupted zip archive header"
    with pytest.raises(CorruptDocumentError, match="corrupt or invalid"):
        service.extract_from_docx(corrupt_bytes)


# --- PDF Extraction Tests ---


def test_extract_from_pdf_mocked_success(service):
    """Verify PDF extraction returns normalized text across pages."""
    mock_page1 = MagicMock()
    mock_page1.extract_text.return_value = "Quarterly Business Review 2026\n\nRevenue grew by 15%."
    mock_page2 = MagicMock()
    mock_page2.extract_text.return_value = "Next Steps:\n- Launch product beta"

    mock_reader = MagicMock()
    mock_reader.pages = [mock_page1, mock_page2]

    with patch("pypdf.PdfReader", return_value=mock_reader):
        extracted = service.extract_from_pdf(b"%PDF-1.4 dummy pdf bytes")
        assert "Quarterly Business Review 2026" in extracted
        assert "Next Steps:" in extracted


def test_extract_from_pdf_scanned_image_only_raises(service):
    """Verify scanned PDF with no digital text raises NoExtractableTextError."""
    mock_page = MagicMock()
    mock_page.extract_text.return_value = ""  # Scanned image with no OCR text

    mock_reader = MagicMock()
    mock_reader.pages = [mock_page]

    with patch("pypdf.PdfReader", return_value=mock_reader):
        with pytest.raises(NoExtractableTextError, match="No extractable text found in the PDF"):
            service.extract_from_pdf(b"%PDF-1.4 dummy pdf")


def test_extract_from_pdf_corrupt(service):
    """Verify damaged PDF bytes raise CorruptDocumentError."""
    with pytest.raises(CorruptDocumentError, match="corrupt or encrypted"):
        service.extract_from_pdf(b"not a valid pdf at all")


# --- Dispatch / extract_document Tests ---


def test_extract_document_dispatcher(service):
    """Verify extract_document routes correctly based on filename extension."""
    txt_content = "Architecture decision records:\nUse FastAPI for speed.".encode("utf-8")
    text, source_type = service.extract_document("adr.txt", txt_content)
    assert source_type == DocumentSourceType.TXT
    assert "Architecture decision records:" in text


def test_extract_document_unsupported_extension(service):
    """Verify unsupported file extensions raise UnsupportedDocumentFormatError."""
    with pytest.raises(UnsupportedDocumentFormatError, match="Unsupported document format"):
        service.extract_document("presentation.pptx", b"dummy")


def test_get_document_extraction_service_singleton():
    """Verify get_document_extraction_service returns singleton instance."""
    s1 = get_document_extraction_service()
    s2 = get_document_extraction_service()
    assert s1 is s2
